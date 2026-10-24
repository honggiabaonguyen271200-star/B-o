#!/usr/bin/env python3
"""Kiểm hàng sẵn từ bảng của các nhà cung cấp (NCC) → data/products.js (thay cho bảng hàng chính gom tay).

Cách dùng:
    python3 scripts/nguon_hang.py <thư mục .xlsx> [--bao-cao bao-cao.md] [--thu]

Trong thư mục có các file .xlsx tải từ Google Sheet của NCC, tên file cho biết kiểu bảng:
    ctv*.xlsx    mỗi dòng một mã, cột "CODE", "GIÁ CTV", các cột size ghi 1 (còn) / 0 (hết)
    hana*.xlsx   nhiều khối cạnh nhau: tên (có mã) → các dòng SIZE / GIÁ / GHI CHÚ (HP, HN… = còn; "đang về" = bỏ)
    trung*.xlsx  dòng "Tên" + dòng cột size + cột "GIÁ", ô size có 1 là còn (giống bảng hàng chính cũ)
    thang*.xlsx  khối "TÊN MẪU-CODE / SIZE / GIÁ / TÌNH TRẠNG", chỉ lấy "SẴN …"
Link các bảng NCC không ghi trong kho mã (kho công khai): nằm trong lịch tự động kiểm hàng.

Chỉ lấy giày đang sẵn (bỏ hàng đang về, quần áo, phụ kiện). Cùng mã ở nhiều NCC: gộp size, mỗi size lấy giá
CTV thấp nhất. Giá bán = giá CTV + PHI (dưới NGUONG) hoặc × (1 + TY_LE) làm tròn lên 10.000đ (chủ shop chốt 24/10).
Tên: data/ten-mau.csv > tên đang có trên web > tên trong bảng NCC > tên tạm theo dòng giày (báo trong báo cáo).
--thu: chỉ in báo cáo, không ghi data/products.js.
"""
import json
import math
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_products as bp  # noqa: E402
import static_pages  # noqa: E402

NGUONG = 3000   # nghìn đồng
PHI = 300       # cộng thêm khi giá CTV dưới ngưỡng
TY_LE = 0.10    # từ ngưỡng trở lên: cộng 10%

SKIP_TABS = re.compile(r"cách làm việc|quy định|note|quần áo|phụ kiện|labubu|pop mart|túi|trang tính", re.I)
TAB_HINTS = [  # tên tạm theo trang (bảng CTV chỉ ghi mã)
    (r"jd1 low", "Jordan", "Air Jordan 1 Low"), (r"jd1|jordan", "Jordan", "Air Jordan"),
    (r"dunk", "Nike", "Nike Dunk"), (r"af1", "Nike", "Nike Air Force 1"),
    (r"airmax|air max|v2k", "Nike", "Nike"), (r"court|cortez|killshot", "Nike", "Nike"),
    (r"new balance", "New Balance", "New Balance"), (r"asics", "Asics", "Asics"),
    (r"adidas|samba|yeezy|adifom", "Adidas", "Adidas"), (r"puma", "Puma", "Puma"),
]


def gia_ban(ctv):
    if ctv < NGUONG:
        return ctv + PHI
    return int(math.ceil(ctv * (1 + TY_LE) / 10.0) * 10)


def txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return bp.clean(v)


def size_label(v):
    """'405' / '40.5' / '40,5' / '9uk-43' / '36.5-4UK' / '41 (rep box)' → '40,5' / '43' / '36,5' / '41'; không phải size → None."""
    s = txt(v).lower().replace(" ", "")
    if not s:
        return None
    for num, unit in re.findall(r"(\d{2,3}(?:[.,]5)?)(uk|us|y)?", s):
        if unit in ("uk", "us"):
            continue
        n = num.replace(".", ",")
        if re.fullmatch(r"\d{3}", n) and n.endswith("5"):  # 405 → 40,5
            n = n[:2] + ",5"
        if re.fullmatch(r"\d{2}(,5)?", n) and 16 <= int(n[:2]) <= 50:
            return n + ("Y" if unit == "y" else "")
    return None


def price_k(v):
    """Giá theo nghìn đồng: 2500 / '2.150.000 đ' / 2150000.0 / '2.290' → 2500 / 2150 / 2150 / 2290."""
    if isinstance(v, (int, float)):
        n = float(v)
        n = n / 1000 if n >= 100000 else n
        return int(n) if n >= 100 else None
    s = txt(v).lower().replace("đ", "").replace("vnd", "").strip()
    m = re.match(r"^(\d{1,3}(?:[.,]\d{3})+|\d+)", s)
    if not m:
        return None
    n = int(re.sub(r"[.,]", "", m.group(1)))
    n = n // 1000 if n >= 100000 else n
    return n if n >= 100 else None


def norm_code(raw):
    s = txt(raw).upper()
    s = re.sub(r"\b([A-Z]{2}\d{4})[ _]+(\d{3})\b", r"\1-\2", s)  # IB0018 101
    s = re.sub(r"\b(\d{6}C?)[_ ]([A-Z]{3,4})\b", r"\1-\2", s)
    s = re.sub(r"\b([A-Z0-9]{5,9})_([A-Z0-9]{2,4})\b", r"\1-\2", s)  # 401698_01
    return s


def code_of(text, only_code=False):
    """Mã trong tên; only_code: ô chỉ ghi mã ('1201A019 107', 'FD9763 101' → nối bằng '-')."""
    s = norm_code(text)
    if only_code:
        s = re.sub(r"^([A-Z0-9]{5,9})\s+(\d{2,3})$", r"\1-\2", s.strip())
    code = bp.find_code(s)
    if not code:
        m = re.search(r"\bP[A-Z0-9]{9,12}\b", s)  # Jeep P740W14555SX
        code = m.group(0) if m else ""
    return code


def read_book(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        if SKIP_TABS.search(ws.title):
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        yield ws.title.strip(), rows


CODE_BRANDS = [  # nhận hãng theo dạng mã khi tên không ghi hãng
    (r"^WRS", "Wilson"), (r"^VN0", "Vans"), (r"^3[A-Z]{2}\d{8}$", "On"), (r"^\d[A-Z0-9]{7,9}-\d{4}$", "Babolat"),
    (r"^\d{6}C?-[A-Z]{3,4}$", "Skechers"), (r"^1\d{3}[A-Z]\d{3}-\d{3}$", "Asics"), (r"^\d{6}-\d{2}$", "Puma"),
    (r"^[A-Z]{2}\d{4}-\d{3}$", "Nike"), (r"^[A-Z]{2}\d{4}$", "Adidas"),
]


def code_brand(code):
    for pat, brand in CODE_BRANDS:
        if re.search(pat, code):
            return brand
    return None


def tab_hint(tab):
    for pat, brand, model in TAB_HINTS:
        if re.search(pat, tab, re.I):
            return brand, model
    return None, None


def item(src, tab, raw, code, sizes, note=None):
    return {"src": src, "tab": tab, "raw": raw, "code": code, "sizes": sizes, "note": note}


# ---------- 4 kiểu bảng ----------

def parse_ctv(path):
    for tab, rows in read_book(path):
        head = None
        for r in rows:
            cells = [txt(c) for c in r]
            if "CODE" in [c.upper() for c in cells]:
                head = cells
                ci = [c.upper() for c in cells].index("CODE")
                pi = next((i for i, c in enumerate(cells) if c.upper().startswith("GIÁ")), ci + 1)
                scols = {i: size_label(c) for i, c in enumerate(cells) if i > pi and size_label(c)}
                continue
            if not head or ci >= len(cells) or not cells[ci]:
                continue
            code = code_of(cells[ci], only_code=True)
            base = price_k(r[pi]) if pi < len(r) else None
            sizes = {}
            for i, s in scols.items():
                v = txt(r[i]) if i < len(r) else ""
                if re.fullmatch(r"[1-9]\d?", v) and base:  # số đôi còn; 0 / trống / "ĐANG VỀ" = không có
                    sizes[s] = base
            if code:
                yield item("ctv", tab, "", code, sizes)


def parse_hana(path):
    for tab, rows in read_book(path):
        blocks = []
        for r in rows[:5]:
            cells = [txt(c).upper() for c in r]
            if "SIZE" in cells:
                blocks = [j for j, c in enumerate(cells) if c == "SIZE" and j > 0]
                break
        for j in blocks:  # tên: cột ngay trước SIZE; giá: SIZE+1; ghi chú: SIZE+2
            cur = None
            for r in rows:
                get = lambda k: txt(r[k]) if k < len(r) else ""  # noqa: E731
                name, size, price, note = get(j - 1), r[j] if j < len(r) else None, r[j + 1] if j + 1 < len(r) else None, get(j + 2)
                if name and name.upper() not in ("ẢNH", "TÊN SẢN PHẨM", "MÃ SP"):
                    if cur:
                        yield cur
                    cur = item("hana", tab, name, code_of(name), {})
                s = size_label(size)
                if cur and s and "đang về" not in note.lower():
                    p = price_k(price)
                    if p:
                        cur["sizes"][s] = min(p, cur["sizes"].get(s, p))
            if cur:
                yield cur


def parse_trung(path):
    for tab, rows in read_book(path):
        scols, pi, i = {}, None, 0
        while i < len(rows):
            cells = [txt(c) for c in rows[i]]
            if cells and cells[0] == "Tên":
                # cột size có thể nằm cùng dòng "Tên" hoặc dòng ngay dưới
                for k in (i, i + 1):
                    if k < len(rows):
                        hc = [txt(c) for c in rows[k]]
                        sc = {n: size_label(c) for n, c in enumerate(hc) if n > 0 and re.fullmatch(r"\d{2}([.,]5)?Y?", c)}
                        if sc:
                            scols = sc
                            pi = next((n for n, c in enumerate(hc) if c.upper().startswith("GIÁ")), None)
                            i = k
                            break
                i += 1
                continue
            if scols and pi is not None and cells and cells[0]:
                raw = cells[0]
                base = price_k(rows[i][pi]) if pi < len(rows[i]) else None
                sizes = {}
                for n, s in scols.items():
                    info = bp.parse_size_cell(rows[i][n] if n < len(rows[i]) else "")
                    v = txt(rows[i][n]) if n < len(rows[i]) else ""
                    if info is not None and re.match(r"^1\b", v):
                        p = info.get("price") or base
                        if p:
                            sizes[s] = p
                extra = " ".join(c for c in cells[pi + 1:] if c and not re.match(r"(?i)^update", c))
                yield item("trung", tab, raw, code_of(raw), sizes, extra or None)
            i += 1


def parse_thang(path):
    for tab, rows in read_book(path):
        names = []
        for r in rows[:6]:
            cells = [txt(c).upper() for c in r]
            if any(c.startswith("TÊN MẪU") for c in cells):
                names = [j for j, c in enumerate(cells) if c.startswith("TÊN MẪU")]
                break
        for j in names:  # tên j, size j+1, giá j+2, tình trạng j+3
            cur = None
            for r in rows:
                get = lambda k: txt(r[k]) if k < len(r) else ""  # noqa: E731
                name = get(j)
                if name and not name.upper().startswith(("TÊN MẪU", "KÉO XUỐNG")):
                    if cur:
                        yield cur
                    cur = item("thang", tab, name, code_of(name), {})
                s = size_label(r[j + 1] if j + 1 < len(r) else None)
                if cur and s and "SẴN" in get(j + 3).upper():
                    p = price_k(r[j + 2] if j + 2 < len(r) else None)
                    if p:
                        cur["sizes"][s] = min(p, cur["sizes"].get(s, p))
            if cur:
                yield cur


PARSERS = {"ctv": parse_ctv, "hana": parse_hana, "trung": parse_trung, "thang": parse_thang}


# ---------- gộp + ghi ----------

def current_products():
    text = bp.OUT.read_text(encoding="utf-8")
    return json.loads(re.search(r"window\.PRODUCTS\s*=\s*(.*);\s*$", text, re.S).group(1))


def clean_name(raw, code):
    name = bp.display_name(re.sub(r"\s+", " ", raw), code) if raw else ""
    name = re.sub(r"^\(?(gs|w)\)?\s+", "", name, flags=re.I)
    name = re.sub(r"\s+(sz|size)\s+khác.*$|\s+có\s+order.*$|\s+-\s*\d{2}(,5)?\s*$", "", name, flags=re.I)  # ghi chú của NCC
    return name.strip(" -/")


def custom_names():
    import csv
    if not bp.NAMES.exists():
        return set()
    with bp.NAMES.open(encoding="utf-8-sig", newline="") as f:
        return {r[0].strip().upper() for r in csv.reader(f) if r and r[0].strip() and r[0] != "Mã"}


def sibling_model(code, old_by):
    """Tên dòng giày của các mẫu cùng mã gốc (FZ2158-600 ← FZ2158-102 'Nike Zoom Vapor Pro 3'): phần chữ đầu giống nhau."""
    base = re.split(r"[-_ ]", code)[0]
    if len(base) < 5:
        return ""
    names = [p["name"] for c, p in old_by.items() if c != code and re.split(r"[-_ ]", c)[0] == base and p.get("name")]
    if not names:
        return ""
    words = [re.sub(r"['\"“”‘’]", "", n).split() for n in names]
    common = []
    for ws in zip(*words):
        if len({w.lower() for w in ws}) != 1:
            break
        common.append(ws[0])
    if len(common) < 2:  # tên khác nhau: lấy tên ngắn nhất, bỏ phần màu sau dấu nháy
        common = min((re.split(r"['\"“‘]", n)[0].split() for n in names), key=len)
    return " ".join(common) if len(common) >= 2 else ""


def merge(items, old):
    by = {}
    errs = []
    for it in items:
        code = it["code"]
        if not code:
            if it["sizes"] and it["raw"]:
                errs.append(f"{it['src']} / {it['tab']}: không thấy mã trong \"{it['raw'][:60]}\"")
            continue
        if bp.APPAREL.search(it["raw"] or ""):
            continue
        rec = by.setdefault(code, {"code": code, "names": [], "tabs": set(), "sizes": {}, "srcs": {}, "notes": set(), "raw": ""})
        if it["raw"]:
            rec["names"].append((it["src"], it["raw"]))
            rec["raw"] += " " + it["raw"]
        rec["tabs"].add(it["tab"])
        if it.get("note"):
            rec["notes"].add(it["note"])
        for s, p in it["sizes"].items():
            if s not in rec["sizes"] or p < rec["sizes"][s]:
                rec["sizes"][s] = p
                rec["srcs"][s] = it["src"]
    old_by = {p["code"]: p for p in old if p.get("code")}
    out, need_name = [], []
    for code, rec in by.items():
        prev = old_by.get(code)
        brand, model = None, None
        for t in rec["tabs"]:
            brand, model = tab_hint(t)
            if brand:
                break
        name = prev["name"] if prev else ""
        if not name:
            for pref in ("trung", "thang", "hana", "ctv"):
                cand = [clean_name(r, code) for s, r in rec["names"] if s == pref]
                cand = [c for c in cand if len(c) > 3]
                if cand:
                    name = cand[0]
                    break
        if not name:
            name = sibling_model(code, old_by)
        if not name:
            name = f"{model or brand or ''} {code}".strip()
            if rec["sizes"] and code not in custom_names():
                need_name.append(code)
        name = re.sub(r"\s+(sz|size)\s+khác.*$|\s+có\s+order.*$|\s+-\s*\d{2}(,5)?\s*$", "", name, flags=re.I).strip()
        if len(name.split()) > 2:  # tên dòng giày + mã (lượt trước): bỏ mã, mã đã hiện riêng trên thẻ
            name = re.sub(r"\s+" + re.escape(code) + r"$", "", name, flags=re.I)
        detected = bp.detect(bp.BRAND_RULES, name + " " + rec["raw"], None)
        brand = (prev or {}).get("brand") or detected or code_brand(code) or brand or "Khác"
        first = brand.split()[0]
        if not re.search(r"\b" + re.escape(first) + r"\b", name, re.I):
            name = f"{brand} {name}"
        gender = (prev or {}).get("gender") or bp.detect(bp.GENDER_RULES, rec["raw"], "unisex")
        sizes = sorted(rec["sizes"].items(), key=lambda kv: (float(kv[0].rstrip("Y").replace(",", ".")), kv[0]))
        retail = {s: gia_ban(p) for s, p in sizes}
        price = min(retail.values()) if retail else (prev or {}).get("price")
        p = {"id": bp.slugify(code), "code": code, "name": name, "brand": brand, "gender": gender, "price": price,
             "sizes": [dict({"s": s}, **({"p": retail[s]} if retail[s] != price else {})) for s, _ in sizes]}
        if rec["notes"]:
            p["note"] = "; ".join(sorted(rec["notes"]))
        p["_srcs"] = rec["srcs"]
        p["_ctv"] = dict(sizes)
        out.append(p)
    return out, need_name, errs


def report(new, old, need_name, errs, counts):
    old_by = {p["code"]: p for p in old}
    new_by = {p["code"]: p for p in new}
    stock = lambda p: {s["s"] for s in p.get("sizes") or []}  # noqa: E731
    came = [c for c, p in new_by.items() if stock(p) and not stock(old_by.get(c, {}))]
    gone = [c for c, p in old_by.items() if stock(p) and not stock(new_by.get(c, {}))]
    changed = [c for c, p in new_by.items() if c in old_by and stock(p) and stock(old_by[c]) and stock(p) != stock(old_by[c])]
    price = [c for c, p in new_by.items() if c in old_by and stock(p) and stock(old_by[c]) and p.get("price") != old_by[c].get("price")]
    L = ["# Báo cáo kiểm hàng", "", "Đọc từ bảng NCC: " + ", ".join(f"{k} {v} mẫu" for k, v in counts.items()), "",
         f"- Còn hàng: {sum(1 for p in new if stock(p))} mẫu (trước {sum(1 for p in old if stock(p))})",
         f"- Mới về / có lại: {len(came)}", f"- Hết hàng (ẩn khỏi web): {len(gone)}",
         f"- Đổi size: {len(changed)}", f"- Đổi giá: {len(price)}", ""]
    def lst(title, codes, fmt):
        if codes:
            L.append(f"## {title} ({len(codes)})")
            L.extend(fmt(c) for c in sorted(codes))
            L.append("")
    lst("Mới về / có lại", came, lambda c: f"- {c} {new_by[c]['name']} · {new_by[c]['price']}k · size {', '.join(sorted(stock(new_by[c])))}")
    lst("Hết hàng", gone, lambda c: f"- {c} {old_by[c]['name']}")
    lst("Đổi giá", price, lambda c: f"- {c} {new_by[c]['name']}: {old_by[c].get('price')}k → {new_by[c]['price']}k")
    lst("Cần đặt tên (bảng chỉ ghi mã)", need_name, lambda c: f"- {c} (đang để: {new_by[c]['name']})")
    lst("Dòng không đọc được mã", errs, lambda e: f"- {e}")
    return "\n".join(L) + "\n"


def main():
    args = sys.argv[1:]
    if not args or args[0].startswith("-"):
        sys.exit(__doc__)
    folder = Path(args[0])
    out_report = Path(args[args.index("--bao-cao") + 1]) if "--bao-cao" in args else None
    items, counts = [], {}
    for f in sorted(folder.glob("*.xlsx")):
        kind = next((k for k in PARSERS if f.name.lower().startswith(k)), None)
        if not kind:
            continue
        got = list(PARSERS[kind](f))
        counts[kind] = sum(1 for g in got if g["sizes"])
        items += got
    old = current_products()
    new, need_name, errs = merge(items, old)
    # mẫu cũ không còn trong bảng NCC nào: giữ trên danh sách nhưng hết size (web tự ẩn)
    have = {p["code"] for p in new}
    for p in old:
        if p.get("code") and p["code"] not in have:
            q = {k: v for k, v in p.items() if k not in ("sizes", "since", "order")}
            q["sizes"] = []
            new.append(q)
    order = {p["code"]: p.get("order", 10 ** 6) for p in old}
    new.sort(key=lambda p: (order.get(p["code"], 10 ** 6), p["brand"], p["name"]))
    text = report(new, old, need_name, errs, counts)
    if out_report:
        out_report.write_text(text, encoding="utf-8")
    print(text)
    if "--thu" in args:
        return
    for i, p in enumerate(new):
        p["order"] = i
        p.pop("_srcs", None)
        p.pop("_ctv", None)
        img = bp.find_image(p["code"])
        if img:
            p["image"] = img
        else:
            p.pop("image", None)
    bp.apply_names(new)
    bp.fix_sizes(new)
    bp.mark_arrivals(new)
    bp.write_products(new)
    static_pages.main()


if __name__ == "__main__":
    main()
